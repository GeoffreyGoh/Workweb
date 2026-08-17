"""One-time import of the client product catalogue from their Excel workbook.

    python import_catalog.py --dry-run          parse and report, write nothing
    python import_catalog.py                    import for real
    python import_catalog.py --file "path.xlsx" a different workbook

The workbook is hand-maintained and not a flat table, so the parser is
deliberately defensive: a malformed row is reported and skipped rather than
being allowed to abort the run.

What the sheets actually look like
----------------------------------
* The first few rows are a legend mapping a cell's *background fill* to a
  stock status. The fills differ per sheet, so the legend is read per sheet.
* Then a header row of repeated "Colour" labels.
* Then the data. Column A holds the product name; every non-empty cell to its
  right is one colour variant, with its status carried by the fill colour.
* A row with a BLANK column A continues the product above it - roughly a fifth
  of all rows. Treating those as separate products would invent products that
  do not exist.

Known quirks handled here, all verified against the 11 August 2026 workbook:
* `Curtain` and `Design Shades` have no legend at all -> status 'unknown'.
* `Curtain` writes decimals with commas ("max drop 3,2m").
* `Curtain` carries fabric codes in the cell ("Sand : 10502").
* `Curtain` has a divider row; everything above it is historic -> 'legacy'.
* Some fills appear in no legend (dark grey FF666666 in particular). Those
  become 'unknown' and are counted in the summary rather than guessed at.

There is no price anywhere in this workbook. Everything is imported at
unit_price = 0 and is_active = False, so nobody can quote an item before the
client supplies a price list.
"""

import argparse
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal

import models
from database import Base, SessionLocal, engine

DEFAULT_FILE = r"C:\Users\Lenovo\Downloads\AS PER 11 August 2026 (1).xlsx"

INCLUDED_SHEETS = [
    "Roller",
    "Design Shades",
    "All Venetian",
    "Curtain",
    "Cellular Shade",
    "Vertical Fabric",
]

# 'Frame Door and Window' is excluded on the product owner's instruction.
# The rest are supplementary or legacy reference sheets.
EXCLUDED_SHEETS = [
    "Frame Door and Window",
    "Note",
    "FABRIC CHART",
    "Venetian Limited stocks",
    "List Fabric has discontinued",
    "Sheet4",
    "Sheet3",
]

# Short code prefix per sheet, used to build the product code.
PREFIX = {
    "Roller": "ROL",
    "Design Shades": "DSH",
    "All Venetian": "VEN",
    "Curtain": "CTN",
    "Cellular Shade": "CEL",
    "Vertical Fabric": "VRT",
}

# An unqualified "max 2.9m" means width on blind sheets and drop on curtains.
UNQUALIFIED_MEANS = defaultdict(lambda: "width", {"Curtain": "height"})

# Legend text -> our stored status.
LEGEND_TEXT = [
    ("OUT OF STOCK", "out_of_stock"),
    ("LIMITED STOCK", "limited_stock"),
    ("IN STOCK", "in_stock"),
    ("WILL DISCONTINUE", "will_discontinue"),
    ("DISCONTINUED", "discontinued"),
    ("NEW FABRIC RANGE", "new_range"),
    ("NEW FABRIC COLOUR", "new_range"),
]

CURTAIN_DIVIDER = re.compile(r"current\s+fabrics\s+available", re.I)

# Column A sometimes holds an annotation or a section banner rather than a
# product: "Note :", "Current Aluminium Venetian stocks available".
SKIP_ROW_RE = re.compile(r"^(note\s*:?\s*$|current\b.*\bavailable\b)", re.I)


# ------------------------------------------------------------------ helpers
def cell_fill(cell):
    """The cell's background colour as an ARGB string, or None."""
    pattern = cell.fill
    if not pattern or pattern.patternType is None:
        return None
    rgb = getattr(pattern.fgColor, "rgb", None)
    if isinstance(rgb, str) and rgb != "00000000":
        return rgb
    return None


def clean(value) -> str:
    """Collapse newlines and runs of whitespace into single spaces."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def read_legend(ws, header_row) -> dict:
    """Map fill colour -> status for THIS sheet, from its own legend rows.

    Only rows ABOVE the "Colour" header count. Data cells sometimes contain a
    status word themselves - Curtain C5 reads "Ice/Artic (max drop 3,2m)
    Limited Stock" - and scanning the whole top of the sheet would read that
    as a legend entry, mapping green to 'limited_stock' and mislabelling every
    in-stock Curtain colour. A sheet with no header row has no legend.
    """
    if not header_row:
        return {}

    legend = {}
    for row in ws.iter_rows(min_row=1, max_row=header_row - 1):
        for cell in row:
            text = clean(cell.value).upper()
            if not text or len(text) > 60:
                continue
            fill = cell_fill(cell)
            if not fill:
                continue
            # The label may carry a qualifier: "STOCK WILL DISCONTINUE
            # (RANGE FABRIC)" or "THESE STOCK WILL DISCONTINUE (...)".
            bare = clean(re.sub(r"\([^)]*\)", " ", text))
            for needle, status in LEGEND_TEXT:
                if bare.endswith(needle):
                    legend.setdefault(fill, status)
                    break
    return legend


def find_header_row(ws):
    """The row of repeated 'Colour' labels; data starts below it."""
    for r in range(1, 16):
        labels = [clean(ws.cell(r, c).value).lower() for c in range(2, 9)]
        if labels.count("colour") >= 2:
            return r
    return None


DIM_RE = re.compile(
    r"max\.?\s*(width|drop|height|length)?\s*[:\-]?\s*(\d+(?:[.,]\d+)?)\s*(m|cm|mm)\b",
    re.I,
)
# "Pristine White (2.7m)" - a bare measurement with no "max" in front of it.
BARE_DIM_RE = re.compile(r"\(\s*(\d+(?:[.,]\d+)?)\s*(m|cm|mm)\s*\)", re.I)
# "Sand : 10502" or "Champagne 53233"
CODE_RE = re.compile(r"(?::\s*|\s+)(\d{4,6})(?!\d)")


def to_cm(number: str, unit: str) -> Decimal:
    """The sheet mixes metres and the odd centimetre; normalise to cm."""
    value = Decimal(number.replace(",", "."))
    unit = unit.lower()
    if unit == "m":
        return value * 100
    if unit == "mm":
        return value / 10
    return value


def parse_color(text: str, sheet: str):
    """Pull a colour name, fabric code, size caps and leftover notes apart.

    Returns (color_name, fabric_code, max_width_cm, max_height_cm, notes).
    """
    original = clean(text)
    working = original
    max_width = max_height = None
    notes = []

    for match in DIM_RE.finditer(original):
        qualifier = (match.group(1) or "").lower()
        centimetres = to_cm(match.group(2), match.group(3))
        if qualifier == "width":
            max_width = centimetres
        elif qualifier in ("drop", "height", "length"):
            max_height = centimetres
        elif UNQUALIFIED_MEANS[sheet] == "width":
            max_width = centimetres
        else:
            max_height = centimetres
        working = working.replace(match.group(0), " ")

    if max_width is None and max_height is None:
        bare = BARE_DIM_RE.search(working)
        if bare:
            centimetres = to_cm(bare.group(1), bare.group(2))
            if UNQUALIFIED_MEANS[sheet] == "width":
                max_width = centimetres
            else:
                max_height = centimetres
            working = working.replace(bare.group(0), " ")

    fabric_code = None
    code = CODE_RE.search(working)
    if code:
        fabric_code = code.group(1)
        working = working.replace(code.group(0), " ")

    # Whatever is left in brackets is a remark, not part of the colour name.
    for remark in re.findall(r"\(([^)]*)\)", working):
        if clean(remark):
            notes.append(clean(remark))
    working = re.sub(r"\([^)]*\)", " ", working)

    working = clean(working).strip(" :-,;")

    # Trailing remarks the sheet writes without brackets, e.g. "Ice/Artic
    # Limited Stock" or "Gray Sheen approx 24 sets".
    trailing = re.search(
        r"\b(approx\.?\s+.*|limited\s+stock.*|appro\b.*|new\s+colour.*)$", working, re.I
    )
    if trailing and trailing.start() > 0:
        notes.append(clean(trailing.group(0)))
        working = clean(working[: trailing.start()])

    return working, fabric_code, max_width, max_height, "; ".join(notes) or None


def make_code(sheet: str, name: str, used: set) -> str:
    """A stable, readable product code: ROL-ASPEN-BO."""
    slug = re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").upper()
    slug = re.sub(r"-{2,}", "-", slug)[:30].strip("-") or "ITEM"
    code = f"{PREFIX.get(sheet, 'GEN')}-{slug}"[:40]
    if code not in used:
        used.add(code)
        return code
    for n in range(2, 100):
        candidate = f"{code[:37]}-{n}"
        if candidate not in used:
            used.add(candidate)
            return candidate
    raise RuntimeError(f"could not allocate a code for {name!r}")


# -------------------------------------------------------------------- parse
def parse_sheet(ws, report):
    """Yield product dicts, each with its list of colour dicts."""
    sheet = ws.title
    header = find_header_row(ws)
    legend = read_legend(ws, header)
    start = (header + 1) if header else 1

    if not legend:
        report["no_legend"].add(sheet)

    # Curtain: everything above the divider row is historic.
    legacy = sheet == "Curtain"
    products = []
    current = None
    previous_name = None
    used_codes = report["used_codes"]

    for r in range(start, ws.max_row + 1):
        name = clean(ws.cell(r, 1).value)

        if name and SKIP_ROW_RE.match(name):
            # "Current Fabrics Available (2025)" also ends the legacy block.
            if CURTAIN_DIVIDER.search(name):
                legacy = False
            report["skipped_rows"].append(f"{sheet}!row {r}: {name[:40]!r}")
            continue

        colour_cells = [
            ws.cell(r, c)
            for c in range(2, ws.max_column + 1)
            if clean(ws.cell(r, c).value)
        ]

        if not name and not colour_cells:
            continue

        # A product's colours can run over many rows. Usually column A is left
        # blank on those rows, but the client also re-types the same name -
        # "BOTANY (127 mm)" appears on 13 consecutive rows. Both mean "more
        # colours for the product above", not a new product.
        if name and name == previous_name and current is not None:
            report["repeated_name_rows"] += 1
            name = ""

        if name:
            # Keep the name exactly as written. Stripping the parenthetical
            # would turn "BRANXTON (BO)" and "BRANXTON (TL)" into two products
            # both called BRANXTON - block out and translucent are different
            # fabrics, and staff must be able to tell them apart.
            descriptions = re.findall(r"\(([^)]*)\)", name)
            current = {
                "name": name,
                "code": make_code(sheet, name, used_codes),
                "category": sheet,
                "description": "; ".join(clean(d) for d in descriptions if clean(d))
                or None,
                "legacy": legacy,
                "row": r,
                "colors": [],
            }
            products.append(current)
            previous_name = name
        elif current is None:
            # Colours before any product name - nothing to attach them to.
            report["orphan_rows"].append(f"{sheet}!row {r}")
            continue
        else:
            report["continuation_rows"] += 1

        for cell in colour_cells:
            try:
                label = clean(cell.value)
                color_name, code, width, height, notes = parse_color(label, sheet)
                if not color_name:
                    report["unnamed_colors"].append(f"{sheet}!{cell.coordinate}")
                    continue

                fill = cell_fill(cell)
                if legacy:
                    status = "legacy"
                elif not legend:
                    status = "unknown"
                    report["no_legend_colors"] += 1
                else:
                    status = legend.get(fill)
                    if status is None:
                        status = "unknown"
                        report["unmapped_fills"][f"{sheet} {fill}"] += 1

                current["colors"].append({
                    "color_name": color_name[:100],
                    "fabric_code": code,
                    "max_width_cm": width,
                    "max_height_cm": height,
                    "stock_status": status,
                    "raw_label": label[:255],
                    "source_ref": f"{sheet}!{cell.coordinate}",
                    "notes": (notes or None) and notes[:255],
                })
            except Exception as exc:  # one bad cell must not stop the import
                report["errors"].append(f"{sheet}!{cell.coordinate}: {exc}")

    return products


# ------------------------------------------------------------------- import
def wipe_transactions(db):
    """Clear demo documents so the placeholder products can be removed.

    Order matters: children before parents.
    """
    counts = {}
    for label, model in [
        ("receipts", models.Receipt),
        ("delivery note items", models.DeliveryNoteItem),
        ("delivery notes", models.DeliveryNote),
        ("purchase order items", models.PurchaseOrderItem),
        ("purchase orders", models.PurchaseOrder),
        ("sales order items", models.SalesOrderItem),
        ("sales orders", models.SalesOrder),
        ("quotation items", models.QuotationItem),
        ("quotations", models.Quotation),
    ]:
        counts[label] = db.query(model).delete(synchronize_session=False)
    db.commit()
    return counts


def run(path: str, dry_run: bool):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required:  pip install openpyxl")

    print(f"Reading {path}")
    workbook = openpyxl.load_workbook(path, data_only=False)

    missing = [s for s in INCLUDED_SHEETS if s not in workbook.sheetnames]
    if missing:
        sys.exit(f"Sheets missing from the workbook: {missing}")

    report = {
        "no_legend": set(),
        "no_legend_colors": 0,
        "unmapped_fills": Counter(),
        "continuation_rows": 0,
        "repeated_name_rows": 0,
        "skipped_rows": [],
        "duplicate_colors": 0,
        "orphan_rows": [],
        "unnamed_colors": [],
        "errors": [],
        "used_codes": set(),
    }

    all_products = []
    for sheet in INCLUDED_SHEETS:
        parsed = parse_sheet(workbook[sheet], report)
        all_products.extend(parsed)
        colours = sum(len(p["colors"]) for p in parsed)
        print(f"  parsed {sheet:<16} {len(parsed):>4} products, {colours:>4} colours")

    skipped = [s for s in workbook.sheetnames if s not in INCLUDED_SHEETS]
    print(f"  skipped {len(skipped)} sheet(s): {', '.join(skipped)}")

    if dry_run:
        print("\n--dry-run: nothing was written.")
        summarise(all_products, report, wiped=None)
        return

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        wiped = wipe_transactions(db)
        removed = db.query(models.Product).delete(synchronize_session=False)
        db.commit()
        print(f"\n  removed {removed} placeholder product(s)")

        for data in all_products:
            product = models.Product(
                code=data["code"],
                name=data["name"][:150],
                category=data["category"],
                description=data["description"],
                price_unit="per_sqm",
                unit_price=Decimal("0.00"),
                # No price in the workbook - keep them out of the quotation
                # picker until the client supplies a price list.
                is_active=False,
            )
            for colour in data["colors"]:
                product.colors.append(models.ProductColor(**colour))
            db.add(product)
        db.commit()

        summarise(all_products, report, wiped)
    finally:
        db.close()


def summarise(products, report, wiped):
    colours = [c for p in products for c in p["colors"]]
    for product in products:
        seen = Counter(c["color_name"].lower() for c in product["colors"])
        report["duplicate_colors"] += sum(n - 1 for n in seen.values() if n > 1)
    by_category = Counter(p["category"] for p in products)
    by_status = Counter(c["stock_status"] for c in colours)

    print("\n" + "=" * 66)
    print("IMPORT SUMMARY")
    print("=" * 66)

    if wiped:
        print("\nCleared to make room (demo transactions):")
        for label, n in wiped.items():
            if n:
                print(f"  {n:>5}  {label}")

    print("\nProducts imported per category:")
    for category, n in by_category.most_common():
        variants = sum(len(p["colors"]) for p in products if p["category"] == category)
        print(f"  {category:<18} {n:>4} products   {variants:>4} colour variants")
    print(f"  {'TOTAL':<18} {len(products):>4} products   {len(colours):>4} colour variants")

    print("\nColour variants by stock status:")
    for status, n in by_status.most_common():
        print(f"  {status:<18} {n:>4}")

    with_code = sum(1 for c in colours if c["fabric_code"])
    with_width = sum(1 for c in colours if c["max_width_cm"])
    with_height = sum(1 for c in colours if c["max_height_cm"])
    print("\nParsed detail:")
    print(f"  fabric codes found        {with_code:>4}")
    print(f"  max width captured        {with_width:>4}")
    print(f"  max drop/height captured  {with_height:>4}")
    print(f"  continuation rows merged  {report['continuation_rows']:>4}"
          "   (blank column A)")
    print(f"  repeated-name rows merged {report['repeated_name_rows']:>4}"
          "   (column A retyped, e.g. BOTANY x13)")
    print(f"  annotation rows skipped   {len(report['skipped_rows']):>4}"
          "   (\"Note :\", section banners)")
    if report["duplicate_colors"]:
        print(f"  duplicate colour names    {report['duplicate_colors']:>4}"
              "   (same colour listed twice on one product)")

    print("\n" + "!" * 66)
    print("NEEDS THE CLIENT'S ATTENTION")
    print("!" * 66)
    print(f"\n  * ALL {len(products)} products imported with unit_price = 0 and")
    print("    is_active = FALSE. The workbook contains no pricing at all.")
    print("    They will not appear in the quotation product picker until a")
    print("    price list is supplied and they are activated.")

    if report["no_legend"]:
        print(f"\n  * {', '.join(sorted(report['no_legend']))} have NO stock legend.")
        print(f"    {report['no_legend_colors']} colour variants were set to 'unknown'")
        print("    rather than guessing status from their fill colour.")

    if report["unmapped_fills"]:
        total = sum(report["unmapped_fills"].values())
        print(f"\n  * {total} colour cells use a fill that is in no legend,")
        print("    also set to 'unknown':")
        for key, n in report["unmapped_fills"].most_common(10):
            print(f"      {key:<28} {n:>4} cells")

    legacy = by_status.get("legacy", 0)
    if legacy:
        print(f"\n  * {legacy} Curtain colour variants sit above the")
        print("    'Current Fabrics Available (2025)' divider and are tagged")
        print("    'legacy', not active stock.")

    for label, items in (("rows with colours but no product name", report["orphan_rows"]),
                         ("cells with no usable colour name", report["unnamed_colors"]),
                         ("cells that failed to parse", report["errors"])):
        if items:
            print(f"\n  * {len(items)} {label}:")
            for item in items[:8]:
                print(f"      {item}")
            if len(items) > 8:
                print(f"      ... and {len(items) - 8} more")

    print()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", default=DEFAULT_FILE)
    parser.add_argument("--dry-run", action="store_true",
                        help="parse and report without touching the database")
    args = parser.parse_args()
    run(args.file, args.dry_run)
